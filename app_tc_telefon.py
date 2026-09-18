# -*- coding: utf-8 -*-
"""
Optik Notlandırma — TC + Telefon Formatı
=========================================
Format:  SOYAD AD   [TC(11, opsiyonel)]   TELEFON(10/11)   CEVAPLAR
Öğrenci NUMARASI YOK. Kimlik = Ad Soyad + TC.

Öğr. Gör. Hatice Tekiş
"""

import streamlit as st
import pandas as pd
import numpy as np
import re
import io
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# ReportLab — PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors as rl_colors
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak
)
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import hashlib

# ── Sayfa ayarları ──────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Optik Notlandırma · TC+Telefon",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── CSS ─────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&family=JetBrains+Mono:wght@400;500&display=swap');

:root {
    --bg:#0d0d14; --bg2:#12111c; --bg3:#1a1927; --border:#1e1d2e; --border2:#2a2840;
    --text:#e2e0f0; --text2:#a09eb8; --text3:#6b6880; --text4:#4a4860;
    --sidebar-bg1:#12111c; --sidebar-bg2:#0f0f1a; --hero-bg1:#13111f; --hero-bg2:#0f1520;
    --input-bg:#12111c; --input-text:#e2e0f0; --bar-bg:#1a1927; --card-info-bg:#12111c;
    --mc-val:#ffffff; --dist-fill:rgba(255,255,255,0.9);
}
@media (prefers-color-scheme: light) {
    :root {
        --bg:#f5f4ff; --bg2:#ffffff; --bg3:#ededf8; --border:#dddaf0; --border2:#c8c4e8;
        --text:#1a1835; --text2:#4a4770; --text3:#7a789a; --text4:#a0a0b8;
        --sidebar-bg1:#f0effe; --sidebar-bg2:#ebe9fc; --hero-bg1:#eeeaff; --hero-bg2:#e8f0ff;
        --input-bg:#ffffff; --input-text:#1a1835; --bar-bg:#e8e6f8; --card-info-bg:#ffffff;
        --mc-val:#1a1835; --dist-fill:rgba(0,0,0,0.75);
    }
}
html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
.stApp { background: var(--bg) !important; color: var(--text) !important; }
[data-testid="stSidebar"] {
    background: linear-gradient(180deg, var(--sidebar-bg1) 0%, var(--sidebar-bg2) 100%) !important;
    border-right: 1px solid var(--border) !important;
}
.hero-wrap {
    background: linear-gradient(135deg, var(--hero-bg1) 0%, var(--hero-bg2) 100%);
    border: 1px solid var(--border); border-radius: 16px; padding: 2.2rem 2.5rem;
    margin-bottom: 1.8rem; position: relative; overflow: hidden;
}
.hero-wrap::before {
    content:''; position:absolute; top:-60px; right:-60px; width:220px; height:220px;
    background: radial-gradient(circle, rgba(124,106,247,0.18) 0%, transparent 70%); pointer-events:none;
}
.hero-title { font-size:2.0rem; font-weight:700; letter-spacing:-0.04em; color:var(--text); margin:0; line-height:1.1; }
.hero-title span { color:#7c6af7; }
.hero-sub { color:var(--text3); font-size:0.9rem; margin-top:0.4rem; font-weight:400; }
.hero-credit {
    display:inline-flex; align-items:center; gap:0.45rem; margin-top:1.1rem;
    background:rgba(124,106,247,0.09); border:1px solid rgba(124,106,247,0.22);
    border-radius:999px; padding:0.28rem 0.9rem; font-size:0.75rem; color:var(--text); font-weight:500;
}
.hero-credit b { color:var(--text); font-weight:700; }
.hero-badge {
    display:inline-block; background:rgba(124,106,247,0.13); border:1px solid rgba(124,106,247,0.28);
    color:#a89ef9; border-radius:999px; padding:0.18rem 0.75rem; font-size:0.72rem; font-weight:600;
    letter-spacing:0.06em; text-transform:uppercase; margin-bottom:0.7rem; font-family:'JetBrains Mono',monospace;
}
.metric-grid { display:grid; grid-template-columns:repeat(5,1fr); gap:0.9rem; margin:1.4rem 0; }
.mc { background:var(--bg2); border:1px solid var(--border); border-radius:12px; padding:1.1rem 1.3rem; position:relative; overflow:hidden; }
.mc::after { content:''; position:absolute; bottom:0; left:0; right:0; height:2px; }
.mc.purple::after { background:linear-gradient(90deg,#7c6af7,#a89ef9); }
.mc.green::after  { background:linear-gradient(90deg,#34d399,#6ee7b7); }
.mc.amber::after  { background:linear-gradient(90deg,#f59e0b,#fbbf24); }
.mc.rose::after   { background:linear-gradient(90deg,#f43f5e,#fb7185); }
.mc.sky::after    { background:linear-gradient(90deg,#0ea5e9,#38bdf8); }
.mc-label { font-size:0.7rem; color:var(--text3); text-transform:uppercase; letter-spacing:0.09em; font-weight:600; }
.mc-val   { font-family:'JetBrains Mono',monospace; font-size:1.9rem; font-weight:500; color:var(--mc-val); margin-top:0.15rem; line-height:1; }
.mc-sub   { font-size:0.72rem; color:var(--text4); margin-top:0.25rem; }
.sec-head { font-size:0.78rem; font-weight:600; letter-spacing:0.1em; text-transform:uppercase; color:var(--text3); margin:1.8rem 0 0.8rem 0; display:flex; align-items:center; gap:0.5rem; }
.sec-head::after { content:''; flex:1; height:1px; background:var(--border); }
.qa-grid { display:grid; grid-template-columns:repeat(auto-fill,minmax(72px,1fr)); gap:0.55rem; margin:0.8rem 0; }
.qa-card { background:var(--bg2); border:1px solid var(--border); border-radius:10px; padding:0.65rem 0.5rem; text-align:center; }
.qa-card .q-no  { font-size:0.65rem; color:var(--text3); font-weight:600; text-transform:uppercase; letter-spacing:0.06em; }
.qa-card .q-pct { font-family:'JetBrains Mono',monospace; font-size:1.1rem; font-weight:600; margin-top:0.1rem; }
.qa-card .q-ans { font-size:0.7rem; color:var(--text3); margin-top:0.15rem; }
.qa-card.easy   { border-color:rgba(52,211,153,0.4); }  .qa-card.easy .q-pct { color:#34d399; }
.qa-card.mid    { border-color:rgba(251,191,36,0.3); }  .qa-card.mid  .q-pct { color:#f59e0b; }
.qa-card.hard   { border-color:rgba(248,113,113,0.4); } .qa-card.hard .q-pct { color:#f87171; }
.step-lbl { font-size:0.68rem; font-weight:700; letter-spacing:0.12em; text-transform:uppercase; color:#7c6af7; margin-bottom:0.4rem; font-family:'JetBrains Mono',monospace; }
.warn-box { background:rgba(251,191,36,0.08); border-left:3px solid #fbbf24; padding:0.6rem 1rem; border-radius:0 8px 8px 0; font-size:0.82rem; color:#d97706; margin:0.25rem 0; font-family:'JetBrains Mono',monospace; }
.ok-box { background:rgba(52,211,153,0.08); border-left:3px solid #34d399; padding:0.7rem 1.1rem; border-radius:0 8px 8px 0; font-size:0.85rem; color:#059669; margin:0.5rem 0; }
.footer-bar { margin-top:3rem; padding:1.2rem 0 0.5rem 0; border-top:1px solid var(--border); text-align:center; font-size:0.75rem; color:var(--text4); line-height:1.8; }
.footer-bar b { color:var(--text3); font-weight:600; }
.stButton > button {
    background:linear-gradient(135deg,#7c6af7,#6b59e8) !important; color:white !important; border:none !important;
    border-radius:9px !important; font-family:'Inter',sans-serif !important; font-weight:600 !important;
    font-size:0.92rem !important; padding:0.65rem 1.8rem !important; transition:all 0.2s !important;
}
.stButton > button:hover { background:linear-gradient(135deg,#8b79f8,#7c6af7) !important; transform:translateY(-1px); box-shadow:0 4px 20px rgba(124,106,247,0.3) !important; }
.stTextInput > div > div > input {
    background:var(--input-bg) !important; border:1px solid var(--border) !important; color:var(--input-text) !important;
    font-family:'JetBrains Mono',monospace !important; border-radius:8px !important; font-size:1rem !important; letter-spacing:0.08em !important;
}
.stTextInput > div > div > input:focus { border-color:#7c6af7 !important; box-shadow:0 0 0 2px rgba(124,106,247,0.2) !important; }
.dist-bar-wrap { margin:1rem 0; }
.dist-row { display:flex; align-items:center; gap:0.7rem; margin:0.3rem 0; }
.dist-lbl { font-size:0.75rem; color:var(--text3); width:70px; text-align:right; font-family:'JetBrains Mono',monospace; }
.dist-bar-bg { flex:1; background:var(--bar-bg); border-radius:4px; height:20px; overflow:hidden; }
.dist-bar-fill { height:100%; border-radius:4px; display:flex; align-items:center; padding-left:8px; }
.dist-bar-fill span { font-size:0.7rem; font-weight:600; color:var(--dist-fill); font-family:'JetBrains Mono',monospace; }
.dist-cnt { font-size:0.73rem; color:var(--text3); width:30px; }
hr { border-color:var(--border) !important; }
.login-wrap { max-width:420px; margin:4rem auto; background:var(--bg2); border:1px solid var(--border); border-radius:18px; padding:2.8rem 2.4rem; position:relative; overflow:hidden; }
.login-wrap::before { content:''; position:absolute; top:-80px; right:-80px; width:260px; height:260px; background:radial-gradient(circle,rgba(124,106,247,0.15) 0%,transparent 70%); pointer-events:none; }
.login-logo { font-size:2.8rem; text-align:center; margin-bottom:0.5rem; }
.login-title { font-size:1.3rem; font-weight:700; letter-spacing:-0.03em; color:var(--text); text-align:center; margin-bottom:0.3rem; }
.login-sub { font-size:0.82rem; color:var(--text3); text-align:center; margin-bottom:1.8rem; }
.kilavuz-wrap { background:var(--bg2); border:1px solid var(--border); border-radius:14px; padding:2rem; }
.kilavuz-step { display:flex; gap:1rem; margin-bottom:1.4rem; align-items:flex-start; }
.kilavuz-num { background:rgba(124,106,247,0.15); border:1px solid rgba(124,106,247,0.3); color:#7c6af7; font-family:'JetBrains Mono',monospace; font-size:0.85rem; font-weight:700; width:32px; height:32px; border-radius:50%; display:flex; align-items:center; justify-content:center; flex-shrink:0; }
.kilavuz-content { flex:1; }
.kilavuz-content b { color:var(--text); display:block; margin-bottom:0.3rem; font-size:0.92rem; }
.kilavuz-content p { color:var(--text3); font-size:0.82rem; line-height:1.7; margin:0; }
.kilavuz-code { background:var(--bg3); border:1px solid var(--border); border-radius:8px; padding:0.7rem 1rem; font-family:'JetBrains Mono',monospace; font-size:0.78rem; color:#7c6af7; margin-top:0.5rem; line-height:1.9; }
[data-testid="stVerticalBlock"],[data-testid="stHorizontalBlock"],[data-testid="column"],[data-testid="stColumn"],[data-testid="block-container"],.element-container,.stMarkdown,section[data-testid="stSidebar"] > div,div[data-testid="stDecoration"] { background:transparent !important; }
[data-testid="stDecoration"] { display:none !important; }
header[data-testid="stHeader"] { background:transparent !important; border-bottom:none !important; }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# CORE FONKSİYONLAR
# ══════════════════════════════════════════════════════════════════════════════

# ── Türkçe karakter onarımı (optik okuyucu / PDF ligatür bozulmaları) ──────────
# Optik okuyucu çıktılarında bazı Türkçe harfler ligatüre dönüşebiliyor.
# Gözlemlenen örnek: "TAﬁKIRAN" -> "TAŞKIRAN". Yeni bozulma görürsen buraya ekle.
TR_FIX = {
    "ﬁ": "Ş",   # fi-ligatürü çoğu optik çıktıda Ş yerine geliyor
    "ﬂ": "Ş",
    "Ý": "İ", "ý": "ı", "Þ": "Ş", "þ": "ş", "Ð": "Ğ", "ð": "ğ",  # cp1254/latin karışması
}
def turkce_duzelt(s):
    for k, v in TR_FIX.items():
        s = s.replace(k, v)
    return s


def _tr_upper(s):
    return (s or "").translate(str.maketrans("iıçğöşü", "İIÇĞÖŞÜ")).upper()


# Türkçe alfabetik sıralama anahtarı (Ad Soyad'a göre sıralamak için)
_TR_ALPHA = " 0123456789AÂBCÇDEFGĞHIİJKLMNOÖPQRSŞTUÜVWXYZ"
def tr_sort_key(s):
    u = _tr_upper(s)
    return [_TR_ALPHA.find(ch) if _TR_ALPHA.find(ch) >= 0 else len(_TR_ALPHA) for ch in u]


def tc_cep_ayir(gruplar):
    """Rakam gruplarından TC (11 hane) ve cep (10/11 hane) ayır."""
    tc, cep = "", ""
    # Birleşik büyük blokları çöz (TC+telefon bitişik işaretlenmişse)
    acik = []
    for g in gruplar:
        if len(g) == 21:                          # TC(11) + cep(10) bitişik
            acik.append(g[:11]); acik.append(g[11:])
        elif len(g) == 22 and g[11:13] == "05":   # TC(11) + 0cep(11) bitişik
            acik.append(g[:11]); acik.append(g[11:])
        else:
            acik.append(g)
    artik = []
    for g in acik:
        if not cep and len(g) == 11 and g.startswith("05"):
            cep = g
        elif not cep and len(g) == 10 and g.startswith("5"):
            cep = "0" + g
        elif not tc and len(g) == 11 and g[0] != "0":
            tc = g
        else:
            artik.append(g)
    for g in artik:  # ikinci geçiş
        if not tc and len(g) == 11 and g[0] != "0":
            tc = g
        elif not cep and len(g) == 10 and g.startswith("5"):
            cep = "0" + g
        elif not cep and len(g) == 11 and g.startswith("05"):
            cep = g
    return tc, cep


def satir_parse_et(satir):
    """
    Satırı parçala:  SOYAD AD  [TC]  TELEFON  CEVAPLAR
    TC opsiyonel (öğrenci işaretlememiş olabilir). Öğrenci numarası YOK.
    """
    satir = turkce_duzelt(satir.strip())
    if not satir:
        return None, None

    m = re.search(r"\d", satir)
    if not m:
        # Rakam hiç yok — sondaki uzun A-E bloğunu cevap kabul et (fallback)
        mm = re.search(r"[A-EXa-ex\s]{10,}$", satir)
        if mm and mm.start() > 0:
            ad = satir[:mm.start()].strip()
            cev = re.sub(r"[^A-EXa-ex ]", "", mm.group()).strip().replace(" ", "0").upper()
            if ad and cev:
                return {"ad_soyad": ad, "tc": "", "cep": "", "cevaplar": cev}, None
        return None, f"Rakam (TC/telefon) bulunamadı → {satir[:50]}"

    ad_soyad = satir[:m.start()].strip()
    if not ad_soyad:
        return None, "Ad Soyad boş"

    rest = satir[m.start():]
    # Cevap bölgesi = sayısal baştan sonraki ilk harften itibaren
    m2 = re.search(r"[A-Za-z]", rest)
    if m2:
        sayisal_bas  = rest[:m2.start()]
        cevap_kuyruk = rest[m2.start():]
    else:
        sayisal_bas  = rest
        cevap_kuyruk = ""

    gruplar = re.findall(r"\d+", sayisal_bas)
    tc, cep = tc_cep_ayir(gruplar)

    # Cevapları temizle: kalan sayı bloklarını at, sadece A-E / X / boşluk bırak
    cevap_kuyruk = re.sub(r"\b\d{3,}\b", "", cevap_kuyruk)
    cevaplar = re.sub(r"[^A-EXa-ex0 ]", "", cevap_kuyruk).strip().replace(" ", "0").upper()
    if not cevaplar:
        return None, "Cevaplar boş"

    return {"ad_soyad": ad_soyad, "tc": tc, "cep": cep, "cevaplar": cevaplar}, None


def ogrenci_puanla(cevaplar, anahtar):
    """Doğru → puan (100/soru); Yanlış/Boş → 0; Anahtar 'X' → iptal (herkese tam puan)."""
    puan = 100 / len(anahtar)
    cevaplar = cevaplar.ljust(len(anahtar), '0')[:len(anahtar)]
    sonuc = []
    for c, a in zip(cevaplar, anahtar):
        if a.upper() == 'X':
            sonuc.append(round(puan, 4))
        elif c not in ('0', ' ', '') and c == a:
            sonuc.append(round(puan, 4))
        else:
            sonuc.append(0)
    return sonuc


def isle(metin, anahtar):
    satirlar = metin.splitlines()
    sonuclar, hatalar = [], []
    for i, satir in enumerate(satirlar, 1):
        veri, hata = satir_parse_et(satir)
        if hata:
            hatalar.append(f"Satır {i}: {hata}")
            continue
        if veri is None:
            continue
        puanlar = ogrenci_puanla(veri["cevaplar"], anahtar)
        kayit = {
            "ad_soyad": veri["ad_soyad"],
            "tc":       veri["tc"],
            "cep":      veri["cep"],
            "cevaplar": veri["cevaplar"],
        }
        for s, p in enumerate(puanlar, 1):
            kayit[f"S{s}"] = p
        kayit["toplam"] = round(sum(puanlar), 2)
        sonuclar.append(kayit)
    return sonuclar, hatalar


def soru_analizi_hesapla(sonuclar, anahtar):
    analiz = []
    n = len(sonuclar)
    for i, dogru_c in enumerate(anahtar, 1):
        dogru = sum(1 for r in sonuclar if r.get(f"S{i}", 0) > 0)
        yanlis = sum(1 for r in sonuclar if r.get(f"S{i}", 0) == 0
                     and len(r.get("cevaplar", "")) >= i
                     and r["cevaplar"][i-1] not in ('0', ' '))
        bos = n - dogru - yanlis
        pct = round(dogru / n * 100, 1) if n else 0
        analiz.append({"soru": i, "anahtar": dogru_c, "dogru": dogru,
                       "yanlis": yanlis, "bos": bos, "pct": pct})
    return analiz


# ══════════════════════════════════════════════════════════════════════════════
# PDF
# ══════════════════════════════════════════════════════════════════════════════

def pdf_olustur(sonuclar, anahtar, dosya_adi="", ders_bilgisi=None, hoca_adi=""):
    if not getattr(pdfmetrics, '_dv_registered', False):
        import os as _os
        _candidates = [
            "/usr/share/fonts/truetype/dejavu",
            "/usr/share/fonts/dejavu",
            _os.path.join(_os.path.dirname(__file__), "fonts"),
        ]
        _FONT_DIR = next((d for d in _candidates
                          if _os.path.exists(_os.path.join(d, "DejaVuSans.ttf"))), None)
        if _FONT_DIR is None:
            raise FileNotFoundError("DejaVuSans.ttf bulunamadı. 'fonts/' klasörüne DejaVu Sans .ttf ekleyin.")
        pdfmetrics.registerFont(TTFont("DV",      f"{_FONT_DIR}/DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("DV-B",    f"{_FONT_DIR}/DejaVuSans-Bold.ttf"))
        pdfmetrics.registerFont(TTFont("DV-I",    f"{_FONT_DIR}/DejaVuSans-Oblique.ttf"))
        pdfmetrics.registerFont(TTFont("DV-BI",   f"{_FONT_DIR}/DejaVuSans-BoldOblique.ttf"))
        pdfmetrics.registerFont(TTFont("DV-Mono", f"{_FONT_DIR}/DejaVuSansMono.ttf"))
        from reportlab.pdfbase.pdfmetrics import registerFontFamily
        registerFontFamily("DV", normal="DV", bold="DV-B", italic="DV-I", boldItalic="DV-BI")
        pdfmetrics._dv_registered = True

    C_PURPLE=rl_colors.HexColor("#7c6af7"); C_DARK=rl_colors.HexColor("#1a1835")
    C_GREY=rl_colors.HexColor("#6b6880");   C_GREY_L=rl_colors.HexColor("#f5f4ff")
    C_GREEN=rl_colors.HexColor("#059669");  C_GREEN_L=rl_colors.HexColor("#d1fae5")
    C_RED=rl_colors.HexColor("#dc2626");    C_RED_L=rl_colors.HexColor("#fee2e2")
    C_AMBER=rl_colors.HexColor("#d97706");  C_AMBER_L=rl_colors.HexColor("#fef3c7")
    C_SKY=rl_colors.HexColor("#0ea5e9");    C_WHITE=rl_colors.white
    C_BORDER=rl_colors.HexColor("#dddaf0"); C_ROW_ALT=rl_colors.HexColor("#f8f7ff")

    toplamlar = [r["toplam"] for r in sonuclar]
    n = len(toplamlar)
    ort=np.mean(toplamlar); std=np.std(toplamlar); medyan=np.median(toplamlar)
    en_yuksek=max(toplamlar); en_dusuk=min(toplamlar)
    gecme_say=len([t for t in toplamlar if t>=50])
    gecme_pct=round(gecme_say/n*100,1) if n else 0
    soru_sayisi=len(anahtar)
    analiz=soru_analizi_hesapla(sonuclar,anahtar)
    tcli = sum(1 for r in sonuclar if r.get("tc"))

    buf=io.BytesIO()
    def ps(name,**kw): return ParagraphStyle(name,**kw)
    s_title=ps("t",fontName="DV-B",fontSize=22,textColor=C_DARK,spaceAfter=4,leading=26)
    s_sub=ps("s",fontName="DV",fontSize=10,textColor=C_GREY,spaceAfter=2)
    s_credit=ps("c",fontName="DV-BI",fontSize=9,textColor=C_PURPLE,spaceAfter=0)
    s_h2=ps("h2",fontName="DV-B",fontSize=13,textColor=C_DARK,spaceBefore=18,spaceAfter=6)
    s_body=ps("b",fontName="DV",fontSize=9,textColor=C_DARK,leading=14)
    s_tc=ps("tc",fontName="DV",fontSize=9,textColor=C_DARK,alignment=TA_CENTER)
    s_tc_bold=ps("tcb",fontName="DV-B",fontSize=9,textColor=C_WHITE,alignment=TA_CENTER)

    PAGE_W,PAGE_H=A4
    def on_page(canvas,doc):
        canvas.saveState()
        canvas.setStrokeColor(C_PURPLE); canvas.setLineWidth(3)
        canvas.line(2*cm,PAGE_H-1.2*cm,PAGE_W-2*cm,PAGE_H-1.2*cm)
        canvas.setFont("DV",7.5); canvas.setFillColor(C_GREY)
        hf=hoca_adi if hoca_adi else "Optik Notlandirma Sistemi"
        canvas.drawString(2*cm,1.2*cm,f"Optik Notlandirma  |  {hf}")
        canvas.drawRightString(PAGE_W-2*cm,1.2*cm,f"Sayfa {doc.page}")
        canvas.restoreState()

    doc=SimpleDocTemplate(buf,pagesize=A4,leftMargin=2*cm,rightMargin=2*cm,
                          topMargin=2*cm,bottomMargin=2*cm,
                          title="Optik Notlandirma Analiz Raporu",
                          author=hoca_adi or "Optik Notlandirma")
    story=[]

    # 1. Başlık
    story.append(Spacer(1,0.5*cm))
    story.append(Paragraph("Optik Notlandirma",s_title))
    story.append(Paragraph("Analiz Raporu",ps("t2",fontName="DV-B",fontSize=16,textColor=C_PURPLE,spaceAfter=6)))
    if dosya_adi: story.append(Paragraph(f"Kaynak: {dosya_adi}",s_sub))
    story.append(Paragraph(f"Cevap Anahtari: <b>{anahtar}</b>  ({soru_sayisi} soru)",s_sub))
    if hoca_adi: story.append(Paragraph(f"Hazirlayan: <b>{hoca_adi}</b>",s_credit))
    if ders_bilgisi:
        db=ders_bilgisi; story.append(Spacer(1,0.3*cm)); rows=[]
        for k,lbl in [("ders_adi","Ders Adi"),("ders_kodu","Ders Kodu"),("bolum","Bolum"),
                      ("donem","Donem"),("akademik_yil","Akademik Yil")]:
            if db.get(k): rows.append([Paragraph(f"<b>{lbl}</b>",s_body),Paragraph(db[k],s_body)])
        if rows:
            dtw=PAGE_W-4*cm
            dt=Table(rows,colWidths=[dtw*0.3,dtw*0.7])
            dt.setStyle(TableStyle([("BACKGROUND",(0,0),(0,-1),C_GREY_L),("FONTNAME",(0,0),(0,-1),"DV-B"),
                ("GRID",(0,0),(-1,-1),0.4,C_BORDER),("TOPPADDING",(0,0),(-1,-1),5),
                ("BOTTOMPADDING",(0,0),(-1,-1),5),("LEFTPADDING",(0,0),(-1,-1),8)]))
            story.append(dt)
    story.append(HRFlowable(width="100%",thickness=1.5,color=C_PURPLE,spaceAfter=16,spaceBefore=10))

    # 2. Sınıf İstatistikleri
    story.append(Paragraph("Sinif Istatistikleri",s_h2))
    def vb(v,col): return ps(f"vb{v}{col}",fontName="DV-B",fontSize=11,textColor=col,alignment=TA_CENTER)
    stat_rows=[
        [Paragraph("<b>Istatistik</b>",s_tc_bold),Paragraph("<b>Deger</b>",s_tc_bold),
         Paragraph("<b>Istatistik</b>",s_tc_bold),Paragraph("<b>Deger</b>",s_tc_bold)],
        [Paragraph("Toplam Ogrenci",s_body),Paragraph(f"<b>{n}</b>",vb(1,C_PURPLE)),
         Paragraph("Gecen (>=50)",s_body),Paragraph(f"<b>{gecme_say}</b>",vb(2,C_GREEN))],
        [Paragraph("Sinif Ortalamasi",s_body),Paragraph(f"<b>{ort:.2f}</b>",vb(3,C_PURPLE)),
         Paragraph("Gecme Orani",s_body),Paragraph(f"<b>%{gecme_pct}</b>",vb(4,C_GREEN))],
        [Paragraph("Standart Sapma",s_body),Paragraph(f"<b>{std:.2f}</b>",vb(5,C_SKY)),
         Paragraph("Kalan",s_body),Paragraph(f"<b>{n-gecme_say}</b>",vb(6,C_RED))],
        [Paragraph("Medyan",s_body),Paragraph(f"<b>{medyan:.2f}</b>",vb(7,C_DARK)),
         Paragraph("En Yuksek",s_body),Paragraph(f"<b>{en_yuksek:.2f}</b>",vb(8,C_GREEN))],
        [Paragraph("TC Isaretleyen",s_body),Paragraph(f"<b>{tcli}</b>",vb(9,C_SKY)),
         Paragraph("En Dusuk",s_body),Paragraph(f"<b>{en_dusuk:.2f}</b>",vb(10,C_RED))],
    ]
    cw=(PAGE_W-4*cm)/4
    st_tbl=Table(stat_rows,colWidths=[cw*1.4,cw*0.6,cw*1.4,cw*0.6])
    st_tbl.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),C_DARK),("TEXTCOLOR",(0,0),(-1,0),C_WHITE),
        ("FONTNAME",(0,0),(-1,0),"DV-B"),("FONTSIZE",(0,0),(-1,0),9),("ALIGN",(0,0),(-1,-1),"CENTER"),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[C_WHITE,C_ROW_ALT]),
        ("GRID",(0,0),(-1,-1),0.5,C_BORDER),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),
        ("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10)]))
    story.append(st_tbl)

    # 3. Geçti/Kaldı
    story.append(Spacer(1,0.4*cm)); story.append(Paragraph("Gecti / Kaldi Ozeti",s_h2))
    gc=[[Paragraph("<b>Durum</b>",s_tc_bold),Paragraph("<b>Sayi</b>",s_tc_bold),
         Paragraph("<b>Oran</b>",s_tc_bold),Paragraph("<b>Puan Esigi</b>",s_tc_bold)],
        [Paragraph("Gecti",ps("g",fontName="DV-B",fontSize=10,textColor=C_GREEN,alignment=TA_CENTER)),
         Paragraph(str(gecme_say),ps("gv",fontName="DV-B",fontSize=14,textColor=C_GREEN,alignment=TA_CENTER)),
         Paragraph(f"%{gecme_pct}",ps("gr",fontName="DV-B",fontSize=12,textColor=C_GREEN,alignment=TA_CENTER)),
         Paragraph(">= 50",s_tc)],
        [Paragraph("Kaldi",ps("k",fontName="DV-B",fontSize=10,textColor=C_RED,alignment=TA_CENTER)),
         Paragraph(str(n-gecme_say),ps("kv",fontName="DV-B",fontSize=14,textColor=C_RED,alignment=TA_CENTER)),
         Paragraph(f"%{round(100-gecme_pct,1)}",ps("kr",fontName="DV-B",fontSize=12,textColor=C_RED,alignment=TA_CENTER)),
         Paragraph("< 50",s_tc)]]
    tw=PAGE_W-4*cm
    gt=Table(gc,colWidths=[tw*0.3,tw*0.2,tw*0.25,tw*0.25])
    gt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),C_DARK),("BACKGROUND",(0,1),(-1,1),C_GREEN_L),
        ("BACKGROUND",(0,2),(-1,2),C_RED_L),("GRID",(0,0),(-1,-1),0.5,C_BORDER),
        ("ALIGN",(0,0),(-1,-1),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    story.append(gt)

    # 4. Puan Dağılımı
    story.append(Paragraph("Puan Dagilimi",s_h2))
    araliklar=[("0-24",0,24,"#f87171"),("25-49",25,49,"#fb923c"),("50-64",50,64,"#fbbf24"),
               ("65-79",65,79,"#34d399"),("80-89",80,89,"#22d3ee"),("90-100",90,100,"#7c6af7")]
    cnts=[sum(1 for t in toplamlar if lo<=t<=hi) for _,lo,hi,_ in araliklar]
    max_cnt=max(cnts) if cnts else 1
    bar_w=PAGE_W-4*cm; bar_h=5.5*cm; lbl_w=1.5*cm; bar_area=bar_w-lbl_w-1.2*cm
    row_h=bar_h/len(araliklar); drw=Drawing(bar_w,bar_h)
    for idx,((lbl,lo,hi,hexc),cnt) in enumerate(zip(araliklar,cnts)):
        y=bar_h-(idx+1)*row_h+row_h*0.15
        drw.add(Rect(lbl_w,y,bar_area,row_h*0.7,fillColor=rl_colors.HexColor("#f0effe"),strokeColor=None))
        fw=(cnt/max_cnt)*bar_area if max_cnt else 0
        if fw>0:
            drw.add(Rect(lbl_w,y,fw,row_h*0.7,fillColor=rl_colors.HexColor(hexc),strokeColor=None))
            drw.add(String(lbl_w+fw+4,y+row_h*0.22,f"%{round(cnt/n*100,1)}  ({cnt})" if n else "",
                           fontName="DV",fontSize=7.5,fillColor=rl_colors.HexColor("#4a4770")))
        drw.add(String(0,y+row_h*0.22,lbl,fontName="DV",fontSize=8,fillColor=rl_colors.HexColor("#6b6880")))
    story.append(drw)

    # 5. Soru Analizi
    story.append(PageBreak()); story.append(Paragraph("Soru Analizi",s_h2))
    story.append(Paragraph("Dogru Yapma Orani:  Kolay >= %70  |  Orta %40-69  |  Zor < %40",
                           ps("leg",fontName="DV",fontSize=8,textColor=C_GREY,spaceAfter=8)))
    qa_rows=[[Paragraph("<b>Soru</b>",s_tc_bold),Paragraph("<b>Cevap</b>",s_tc_bold),
              Paragraph("<b>Dogru</b>",s_tc_bold),Paragraph("<b>Yanlis</b>",s_tc_bold),
              Paragraph("<b>Bos</b>",s_tc_bold),Paragraph("<b>%Dogru</b>",s_tc_bold),
              Paragraph("<b>Zorluk</b>",s_tc_bold),Paragraph("<b>Durum</b>",s_tc_bold)]]
    for a in analiz:
        if a["pct"]>=70:   zt,zf=("Kolay",C_GREEN)
        elif a["pct"]>=40: zt,zf=("Orta",C_AMBER)
        else:              zt,zf=("Zor",C_RED)
        bar="█"*int(a["pct"]/10)+"░"*(10-int(a["pct"]/10))
        qa_rows.append([Paragraph(f"S{a['soru']}",s_tc),
            Paragraph(f"<b>{a['anahtar']}</b>",ps("an",fontName="DV-B",fontSize=9,textColor=C_PURPLE,alignment=TA_CENTER)),
            Paragraph(str(a["dogru"]),ps("d",fontName="DV-B",fontSize=9,textColor=C_GREEN,alignment=TA_CENTER)),
            Paragraph(str(a["yanlis"]),ps("y",fontName="DV-B",fontSize=9,textColor=C_RED,alignment=TA_CENTER)),
            Paragraph(str(a["bos"]),ps("bo",fontName="DV",fontSize=9,textColor=C_GREY,alignment=TA_CENTER)),
            Paragraph(f"{a['pct']:.1f}%",ps("p",fontName="DV-B",fontSize=9,textColor=zf,alignment=TA_CENTER)),
            Paragraph(zt,ps(f"z{a['soru']}",fontName="DV-B",fontSize=8.5,textColor=zf,alignment=TA_CENTER)),
            Paragraph(bar,ps("bar",fontName="DV-Mono",fontSize=7,textColor=zf,alignment=TA_CENTER))])
    tw=PAGE_W-4*cm
    qt=Table(qa_rows,colWidths=[tw*0.08,tw*0.10,tw*0.09,tw*0.09,tw*0.07,tw*0.10,tw*0.10,tw*0.37])
    qs=[("BACKGROUND",(0,0),(-1,0),C_DARK),("GRID",(0,0),(-1,-1),0.4,C_BORDER),
        ("ALIGN",(0,0),(-1,-1),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),("FONTSIZE",(0,0),(-1,0),8.5)]
    for ri,a in enumerate(analiz,start=1):
        bg=C_GREEN_L if a["pct"]>=70 else (C_AMBER_L if a["pct"]>=40 else C_RED_L)
        qs.append(("BACKGROUND",(6,ri),(7,ri),bg))
        if ri%2==0:
            for col in range(6): qs.append(("BACKGROUND",(col,ri),(col,ri),C_ROW_ALT))
    qt.setStyle(TableStyle(qs)); story.append(qt)

    # 6. Öğrenci Listesi (Ad Soyad + TC)
    story.append(PageBreak()); story.append(Paragraph("Ogrenci Listesi",s_h2))
    sirali=sorted(sonuclar,key=lambda r:tr_sort_key(r["ad_soyad"]))
    stu_rows=[[Paragraph("<b>#</b>",s_tc_bold),Paragraph("<b>Ad Soyad</b>",s_tc_bold),
               Paragraph("<b>TC No</b>",s_tc_bold),Paragraph("<b>Toplam</b>",s_tc_bold),
               Paragraph("<b>Durum</b>",s_tc_bold)]]
    for idx,r in enumerate(sirali,1):
        gecti=r["toplam"]>=50
        fg=C_GREEN if gecti else C_RED
        stu_rows.append([
            Paragraph(str(idx),ps(f"rn{idx}",fontName="DV",fontSize=8,textColor=C_PURPLE,alignment=TA_CENTER)),
            Paragraph(r["ad_soyad"],ps(f"ad{idx}",fontName="DV",fontSize=8.5,textColor=C_DARK)),
            Paragraph(r.get("tc","") or "—",ps(f"tc{idx}",fontName="DV",fontSize=8.5,textColor=C_DARK,alignment=TA_CENTER)),
            Paragraph(f"<b>{r['toplam']:.2f}</b>",ps(f"tp{idx}",fontName="DV-B",fontSize=9,textColor=fg,alignment=TA_CENTER)),
            Paragraph("Gecti" if gecti else "Kaldi",ps(f"dt{idx}",fontName="DV-B",fontSize=8,textColor=fg,alignment=TA_CENTER))])
    tw=PAGE_W-4*cm
    stbl=Table(stu_rows,colWidths=[tw*0.07,tw*0.42,tw*0.22,tw*0.15,tw*0.14],repeatRows=1)
    ss=[("BACKGROUND",(0,0),(-1,0),C_DARK),("GRID",(0,0),(-1,-1),0.4,C_BORDER),
        ("ALIGN",(0,0),(-1,-1),"CENTER"),("ALIGN",(1,1),(1,-1),"LEFT"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),("LEFTPADDING",(1,1),(1,-1),6)]
    for idx,r in enumerate(sirali,1):
        gecti=r["toplam"]>=50
        rb=C_WHITE if idx%2 else C_ROW_ALT
        for col in range(4): ss.append(("BACKGROUND",(col,idx),(col,idx),rb))
        ss.append(("BACKGROUND",(4,idx),(4,idx),C_GREEN_L if gecti else C_RED_L))
    stbl.setStyle(TableStyle(ss)); story.append(stbl)

    doc.build(story,onFirstPage=on_page,onLaterPages=on_page)
    buf.seek(0); return buf


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL
# ══════════════════════════════════════════════════════════════════════════════

def excel_olustur(sonuclar, anahtar):
    soru_sayisi=len(anahtar); puan_per_s=round(100/soru_sayisi,4)
    sirali=sorted(sonuclar,key=lambda r:tr_sort_key(r["ad_soyad"]))

    # Sheet 1: Not Girişi (Ad Soyad + TC + her soru puanı)
    rows=[]
    for idx,r in enumerate(sirali,1):
        row={"#":idx,"Adı Soyadı":r["ad_soyad"],"TC No":r.get("tc","")}
        for i in range(1,soru_sayisi+1): row[f"Soru {i}"]=r[f"S{i}"]
        rows.append(row)
    df_not=pd.DataFrame(rows)

    # Sheet 2: Detaylı Liste
    rows=[]
    for idx,r in enumerate(sirali,1):
        row={"#":idx,"Adı Soyadı":r["ad_soyad"],"TC No":r.get("tc",""),
             "Cep Telefonu":r.get("cep",""),"Toplam Puan":r["toplam"],
             "Başarı Durumu":"✓ Geçti" if r["toplam"]>=50 else "✗ Kaldı"}
        for i in range(1,soru_sayisi+1): row[f"S{i}"]=r[f"S{i}"]
        rows.append(row)
    df_detay=pd.DataFrame(rows)

    # Sheet 3: Özet
    df_ozet=pd.DataFrame([{"#":i,"Adı Soyadı":r["ad_soyad"],"TC No":r.get("tc",""),
        "Cep Telefonu":r.get("cep",""),"Toplam Puan":r["toplam"],
        "Başarı Durumu":"Geçti" if r["toplam"]>=50 else "Kaldı"} for i,r in enumerate(sirali,1)])

    # Sheet 4: Soru Analizi
    analiz=soru_analizi_hesapla(sonuclar,anahtar)
    df_analiz=pd.DataFrame([{"Soru No":a["soru"],"Doğru Cevap":a["anahtar"],"Doğru Yapan":a["dogru"],
        "Yanlış Yapan":a["yanlis"],"Boş Bırakan":a["bos"],"Doğru Oranı (%)":a["pct"],
        "Zorluk":"Kolay" if a["pct"]>=70 else ("Orta" if a["pct"]>=40 else "Zor")} for a in analiz])

    # Sheet 5: İstatistik
    toplamlar=[r["toplam"] for r in sonuclar]; gecen=[t for t in toplamlar if t>=50]
    tcli=sum(1 for r in sonuclar if r.get("tc")); cepli=sum(1 for r in sonuclar if r.get("cep"))
    df_istat=pd.DataFrame([
        {"İstatistik":"Toplam Öğrenci","Değer":len(sonuclar)},
        {"İstatistik":"Geçen Öğrenci","Değer":len(gecen)},
        {"İstatistik":"Kalan Öğrenci","Değer":len(sonuclar)-len(gecen)},
        {"İstatistik":"Geçme Oranı (%)","Değer":round(len(gecen)/len(sonuclar)*100,1) if sonuclar else 0},
        {"İstatistik":"Sınıf Ortalaması","Değer":round(np.mean(toplamlar),2) if toplamlar else 0},
        {"İstatistik":"Standart Sapma","Değer":round(np.std(toplamlar),2) if toplamlar else 0},
        {"İstatistik":"En Yüksek Puan","Değer":max(toplamlar) if toplamlar else 0},
        {"İstatistik":"En Düşük Puan","Değer":min(toplamlar) if toplamlar else 0},
        {"İstatistik":"Medyan Puan","Değer":round(np.median(toplamlar),2) if toplamlar else 0},
        {"İstatistik":"TC İşaretleyen","Değer":tcli},
        {"İstatistik":"Cep Tespit Edilen","Değer":cepli},
        {"İstatistik":"Soru Sayısı","Değer":soru_sayisi},
        {"İstatistik":"Soru Başına Puan","Değer":puan_per_s},
        {"İstatistik":"Cevap Anahtarı","Değer":anahtar},
    ])

    # Sheet 6: Ham TXT
    df_ham=pd.DataFrame([{"#":i,"Adı Soyadı":r["ad_soyad"],"TC No":r.get("tc",""),
        "Cep Telefonu":r.get("cep",""),"Ham Cevaplar":r.get("cevaplar","")} for i,r in enumerate(sirali,1)])

    buf=io.BytesIO()
    with pd.ExcelWriter(buf,engine="openpyxl") as w:
        df_not.to_excel(w,sheet_name="Not Girişi",index=False)
        df_detay.to_excel(w,sheet_name="Detaylı Liste",index=False)
        df_ozet.to_excel(w,sheet_name="Özet",index=False)
        df_analiz.to_excel(w,sheet_name="Soru Analizi",index=False)
        df_istat.to_excel(w,sheet_name="İstatistik",index=False)
        df_ham.to_excel(w,sheet_name="Ham TXT",index=False)
    buf.seek(0)
    wb=load_workbook(buf); _excel_stillendir(wb)
    buf2=io.BytesIO(); wb.save(buf2); buf2.seek(0); return buf2


def _excel_stillendir(wb):
    C_HEADER_BG="1E1B35"; C_HEADER_FG="FFFFFF"; C_ALT1="F8F7FF"; C_ALT2="FFFFFF"
    C_ACCENT="7C6AF7"; C_GREEN_BG="EAFAF1"; C_RED_BG="FEF2F2"; C_GREEN_FG="1A6B3C"; C_RED_FG="9B1C1C"
    C_EASY="D1FAE5"; C_MID="FEF3C7"; C_HARD="FEE2E2"; C_NUM_BG="EEF2FF"
    thin=Border(left=Side(style="thin",color="E0DEEE"),right=Side(style="thin",color="E0DEEE"),
                top=Side(style="thin",color="E0DEEE"),bottom=Side(style="thin",color="E0DEEE"))

    def header_row(ws,cols):
        for ci,cn in enumerate(cols,1):
            c=ws.cell(row=1,column=ci,value=cn)
            c.font=Font(bold=True,color=C_HEADER_FG,size=10,name="Calibri")
            c.fill=PatternFill("solid",fgColor=C_HEADER_BG)
            c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            c.border=Border(left=Side(style="thin",color="4A4770"),right=Side(style="thin",color="4A4770"),
                            bottom=Side(style="medium",color=C_ACCENT))
        ws.row_dimensions[1].height=28
    def alt_fill(ri): return PatternFill("solid",fgColor=C_ALT1 if ri%2==0 else C_ALT2)
    def auto_width(ws,min_w=8,max_w=30):
        for col in ws.columns:
            length=max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width=max(min_w,min(length+2,max_w))

    # Sheet 1: Not Girişi
    ws=wb["Not Girişi"]; cn=[c[0].value for c in ws.iter_cols(1,ws.max_column,1,1)]; header_row(ws,cn)
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
        for ci,cell in enumerate(row):
            cell.fill=alt_fill(row[0].row); cell.border=thin
            cell.alignment=Alignment(horizontal="center" if ci!=1 else "left",vertical="center")
            cell.font=Font(name="Calibri",size=10,color="1A1835")
            if ci==0:
                cell.fill=PatternFill("solid",fgColor=C_NUM_BG); cell.font=Font(name="Calibri",size=9,color="7C6AF7",bold=True)
            elif ci==1:
                cell.font=Font(name="Calibri",size=10,bold=True,color="1A1835")
            elif ci>=3:
                if cell.value and float(cell.value)>0: cell.font=Font(name="Calibri",size=10,color="1A6B3C",bold=True)
                else: cell.font=Font(name="Calibri",size=10,color="9B1C1C")
    ws.freeze_panes="D2"; auto_width(ws); ws.column_dimensions["A"].width=5

    # Sheet 2 & 3: Detaylı / Özet
    for sheet in ["Detaylı Liste","Özet"]:
        ws=wb[sheet]; cn=[c[0].value for c in ws.iter_cols(1,ws.max_column,1,1)]; header_row(ws,cn)
        try:
            tcol=cn.index("Toplam Puan")+1; bcol=cn.index("Başarı Durumu")+1
        except ValueError:
            tcol=bcol=None
        for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
            gecti=None
            if tcol:
                pv=row[tcol-1].value; gecti=pv is not None and float(pv)>=50
            for ci,cell in enumerate(row):
                cell.border=thin; cell.alignment=Alignment(horizontal="center" if ci!=1 else "left",vertical="center")
                if ci==0:
                    cell.fill=PatternFill("solid",fgColor=C_NUM_BG); cell.font=Font(name="Calibri",size=9,color="7C6AF7",bold=True)
                elif tcol and ci==tcol-1:
                    cell.fill=PatternFill("solid",fgColor=C_GREEN_BG if gecti else C_RED_BG)
                    cell.font=Font(name="Calibri",size=10,bold=True,color=C_GREEN_FG if gecti else C_RED_FG)
                elif bcol and ci==bcol-1:
                    cell.fill=PatternFill("solid",fgColor=C_GREEN_BG if gecti else C_RED_BG)
                    cell.font=Font(name="Calibri",size=10,bold=True,color=C_GREEN_FG if gecti else C_RED_FG)
                else:
                    cell.fill=alt_fill(row[0].row); cell.font=Font(name="Calibri",size=10,color="1A1835")
        if sheet=="Detaylı Liste": ws.freeze_panes="D2"
        auto_width(ws); ws.column_dimensions["A"].width=5

    # Sheet 4: Soru Analizi
    ws=wb["Soru Analizi"]; cn=[c[0].value for c in ws.iter_cols(1,ws.max_column,1,1)]; header_row(ws,cn)
    try:
        zcol=cn.index("Zorluk")+1; ocol=cn.index("Doğru Oranı (%)")+1
    except ValueError:
        zcol=ocol=None
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
        zv=row[zcol-1].value if zcol else None
        bg=C_EASY if zv=="Kolay" else (C_MID if zv=="Orta" else C_HARD)
        fg=C_GREEN_FG if zv=="Kolay" else ("92400E" if zv=="Orta" else C_RED_FG)
        for ci,cell in enumerate(row):
            cell.border=thin; cell.alignment=Alignment(horizontal="center",vertical="center")
            cell.fill=alt_fill(row[0].row); cell.font=Font(name="Calibri",size=10,color="1A1835")
            if zcol and ci==zcol-1:
                cell.fill=PatternFill("solid",fgColor=bg); cell.font=Font(name="Calibri",size=10,bold=True,color=fg)
            if ocol and ci==ocol-1 and cell.value is not None:
                ov=float(cell.value)
                col=C_GREEN_FG if ov>=70 else ("92400E" if ov>=40 else C_RED_FG)
                cell.font=Font(name="Calibri",size=10,bold=True,color=col)
    auto_width(ws)
    if ocol:
        cl=get_column_letter(ocol)
        ws.conditional_formatting.add(f"{cl}2:{cl}{ws.max_row}",
            ColorScaleRule(start_type="num",start_value=0,start_color="FEE2E2",
                           mid_type="num",mid_value=50,mid_color="FEF3C7",
                           end_type="num",end_value=100,end_color="D1FAE5"))

    # Sheet 5: İstatistik
    ws=wb["İstatistik"]; header_row(ws,["İstatistik","Değer"])
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
        for ci,cell in enumerate(row):
            cell.border=thin; cell.alignment=Alignment(horizontal="left" if ci==0 else "right",vertical="center")
            if ci==0:
                cell.fill=PatternFill("solid",fgColor=C_ALT1); cell.font=Font(name="Calibri",size=10,bold=True,color="3B3860")
            else:
                cell.fill=PatternFill("solid",fgColor=C_ALT2); cell.font=Font(name="Calibri",size=10,color="1A1835")
    auto_width(ws)

    # Sheet 6: Ham TXT
    ws=wb["Ham TXT"]; cn=[c[0].value for c in ws.iter_cols(1,ws.max_column,1,1)]; header_row(ws,cn)
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
        for ci,cell in enumerate(row):
            cell.border=thin; cell.alignment=Alignment(horizontal="center" if ci not in (1,4) else "left",vertical="center")
            cell.fill=alt_fill(row[0].row)
            cell.font=Font(name="Calibri",size=10,color="7C6AF7" if ci==0 else "1A1835",bold=(ci==0))
            if ci==4: cell.font=Font(name="Consolas",size=9,color="4A4770")
    auto_width(ws,max_w=60); ws.column_dimensions["A"].width=5


# ══════════════════════════════════════════════════════════════════════════════
# ARAYÜZ
# ══════════════════════════════════════════════════════════════════════════════

KULLANICI_ADI="tekis-jet"
SIFRE_HASH=hashlib.sha256("tekis2024".encode()).hexdigest()
def sifre_dogrula(s): return hashlib.sha256(s.encode()).hexdigest()==SIFRE_HASH

if "giris_yapildi" not in st.session_state: st.session_state.giris_yapildi=False
if "hoca_adi" not in st.session_state: st.session_state.hoca_adi=""

if not st.session_state.giris_yapildi:
    st.markdown("""
    <div class="login-wrap">
        <div class="login-logo">📊</div>
        <div class="login-title">Optik Notlandırma</div>
        <div class="login-sub">TC + Telefon formatı · Sisteme giriş yapın</div>
    </div>""", unsafe_allow_html=True)
    _,cm_,_=st.columns([1,2,1])
    with cm_:
        hoca=st.text_input("Adınız Soyadınız",placeholder="Öğr. Gör. Ad Soyad")
        ku=st.text_input("Kullanıcı Adı",placeholder="kullanıcı adı")
        si=st.text_input("Şifre",type="password",placeholder="••••••••")
        if st.button("Giriş Yap",use_container_width=True):
            if ku==KULLANICI_ADI and sifre_dogrula(si):
                st.session_state.giris_yapildi=True
                st.session_state.hoca_adi=hoca.strip() or "Öğretim Görevlisi"; st.rerun()
            else:
                st.error("❌ Kullanıcı adı veya şifre hatalı.")
    st.stop()

st.markdown(f"""
<div class="hero-wrap">
  <div class="hero-badge">TC + Telefon · Otomatik Notlandırma</div>
  <h1 class="hero-title">Optik <span>Notlandırma</span> Sistemi</h1>
  <p class="hero-sub">SOYAD AD · TC (opsiyonel) · Telefon · Cevaplar → Excel + PDF</p>
  <div class="hero-credit">✦ &nbsp;<b>{st.session_state.hoca_adi}</b></div>
</div>""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown('<p class="step-lbl">① Cevap Anahtarı</p>',unsafe_allow_html=True)
    anahtar_input=st.text_input("Cevap anahtarı",placeholder="örn: ABCDABCDABCD",
                                label_visibility="collapsed").strip().upper()
    if anahtar_input:
        st.markdown(f'<div style="font-family:JetBrains Mono,monospace;font-size:0.75rem;color:#7c6af7;margin-top:0.3rem;">'
                    f'{len(anahtar_input)} soru · {round(100/len(anahtar_input),2)} pt/soru</div>',unsafe_allow_html=True)
    st.markdown('<p class="step-lbl" style="margin-top:1.5rem;">② TXT Dosyası</p>',unsafe_allow_html=True)
    yuklenen=st.file_uploader("Optik okuyucu TXT",type=["txt"],label_visibility="collapsed")
    st.markdown('<div style="margin-top:1.8rem"></div>',unsafe_allow_html=True)
    isle_btn=st.button("▶  Notlandır",use_container_width=True)
    st.markdown("---")
    st.markdown('<p class="step-lbl" style="margin-top:0;">③ Ders Bilgileri</p>',unsafe_allow_html=True)
    ders_adi=st.text_input("Ders Adı",placeholder="örn: Matematik I",key="ders_adi")
    ders_kodu=st.text_input("Ders Kodu",placeholder="örn: MAT101",key="ders_kodu")
    bolum=st.text_input("Bölüm",placeholder="örn: Bilgisayar Müh.",key="bolum")
    donem=st.selectbox("Dönem",["Güz","Bahar","Yaz"],key="donem")
    akademik_yil=st.text_input("Akademik Yıl",placeholder="2025-2026",key="akademik_yil")
    st.markdown("---")
    st.markdown("""
    <div style="font-size:0.72rem;color:#4a4860;line-height:1.8;">
    <b style="color:#6b6880;">Beklenen Satır Formatı</b><br>
    <span style="color:#7c6af7;">SOYAD AD</span> · TC(11, ops.) · Tel(10/11) · Cevaplar<br><br>
    TC işaretlenmemişse otomatik atlanır.<br>
    Telefon 5… veya 05… olabilir.
    </div>""",unsafe_allow_html=True)
    st.markdown("---")
    if st.button("🚪 Çıkış Yap",use_container_width=True):
        st.session_state.giris_yapildi=False; st.session_state.hoca_adi=""; st.rerun()

if not isle_btn:
    t1,t2,t3=st.tabs(["📌 TXT Formatı","📊 Excel Çıktısı","📖 Kullanım Kılavuzu"])
    with t1:
        st.markdown("""
        <div class="kilavuz-wrap">
        <div style="font-size:0.7rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:var(--text3);margin-bottom:1rem;">Desteklenen TXT Formatı</div>
        <div style="font-family:'JetBrains Mono',monospace;font-size:0.8rem;color:var(--text3);line-height:2.3;">
            <span style="color:#7c6af7;font-weight:700;">Temel</span> — Soyad Ad + TC + Telefon + Cevaplar<br>
            <span style="background:var(--bg3);padding:0.1rem 0.5rem;border-radius:4px;color:var(--text);">TAŞKIRAN NAZLI &nbsp;<b style="color:#a89ef9;">58249498256</b>&nbsp;<b style="color:#34d399;">5522656320</b>&nbsp;CDBCDCBBBE…</span><br><br>
            <span style="color:#7c6af7;font-weight:700;">TC yok</span> — Öğrenci TC işaretlememiş<br>
            <span style="background:var(--bg3);padding:0.1rem 0.5rem;border-radius:4px;color:var(--text);">YILMAZ MEHMET &nbsp;<b style="color:#34d399;">5551234567</b>&nbsp;ABCDEABCDE…</span><br><br>
            <span style="color:#7c6af7;font-weight:700;">Bitişik</span> — TC + Telefon bitişik işaretlenmiş<br>
            <span style="background:var(--bg3);padding:0.1rem 0.5rem;border-radius:4px;color:var(--text);">DEMİR ALİ &nbsp;<b>123456789015551112233</b>&nbsp;ABCDE…</span><br><br>
            <span style="color:var(--text3);font-size:0.75rem;">Telefon 5XXXXXXXXX (10 hane) ya da 05XXXXXXXXX (11 hane) olabilir · Boş bırakılan sorular 0 sayılır · Anahtarda X = iptal soru (herkese tam puan)</span>
        </div>
        </div>""",unsafe_allow_html=True)
    with t2:
        st.markdown("""
        <div class="kilavuz-wrap">
        <div style="font-size:0.7rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:var(--text3);margin-bottom:1rem;">Excel Çıktısı — 6 Sheet</div>
        <div style="font-size:0.85rem;color:var(--text3);line-height:2.5;">
            <span style="color:#7c6af7;">①</span> <b style="color:#7c6af7;">Not Girişi</b> — Ad Soyad + TC + her soru puanı<br>
            <span style="color:#7c6af7;">②</span> <b style="color:#7c6af7;">Detaylı Liste</b> — TC · Cep · Toplam · Durum · tüm soru puanları<br>
            <span style="color:#34d399;">③</span> <b style="color:#34d399;">Özet</b> — Ad Soyad · TC · Toplam · Geçti/Kaldı<br>
            <span style="color:#f59e0b;">④</span> <b style="color:#f59e0b;">Soru Analizi</b> — doğru oranı · zorluk · renk kodlu<br>
            <span style="color:#f59e0b;">⑤</span> <b style="color:#f59e0b;">İstatistik</b> — ortalama · std · medyan · TC/cep sayısı<br>
            <span style="color:var(--text3);">⑥</span> <b style="color:var(--text2);">Ham TXT</b> — ham cevap dizisi (hata ayıklama)<br><br>
            <span style="color:#34d399;">↑</span> Tüm listeler Ad Soyad'a göre Türkçe alfabetik sıralanır
        </div>
        </div>""",unsafe_allow_html=True)
    with t3:
        st.markdown("""
        <div class="kilavuz-wrap">
        <div style="font-size:0.7rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:var(--text3);margin-bottom:1.2rem;">Adım Adım Kullanım</div>
        <div class="kilavuz-step"><div class="kilavuz-num">1</div><div class="kilavuz-content"><b>Cevap Anahtarını Girin</b><p>Sol panele sınav cevaplarını büyük harfle girin (A-E). İptal soru için o pozisyona X yazın.</p><div class="kilavuz-code">örn: ABCDEABCDE → 10 soru · X = iptal (herkese tam puan)</div></div></div>
        <div class="kilavuz-step"><div class="kilavuz-num">2</div><div class="kilavuz-content"><b>TXT Dosyasını Yükleyin</b><p>Optik okuyucudan alınan .txt dosyasını yükleyin. UTF-8, CP1254 ve Latin-1 encoding otomatik denenir.</p></div></div>
        <div class="kilavuz-step"><div class="kilavuz-num">3</div><div class="kilavuz-content"><b>Ders Bilgileri (opsiyonel)</b><p>Ders adı, kodu, bölüm, dönem ve yıl PDF rapora eklenir.</p></div></div>
        <div class="kilavuz-step"><div class="kilavuz-num">4</div><div class="kilavuz-content"><b>Notlandır</b><p>Her satır ayrıştırılır: Ad Soyad, TC (varsa), telefon, cevaplar. 100 üzerinden puanlanır.</p></div></div>
        <div class="kilavuz-step"><div class="kilavuz-num">5</div><div class="kilavuz-content"><b>İndirin</b><p><b style="color:#34d399;">Excel (.xlsx)</b> — 6 sayfa · <b style="color:#f87171;">PDF (.pdf)</b> — analiz raporu</p></div></div>
        <div class="kilavuz-step"><div class="kilavuz-num">!</div><div class="kilavuz-content"><b>Sık Sorunlar</b><p>• <b>Türkçe karakter bozuk (ör. TAﬁKIRAN):</b> kod içindeki <code>TR_FIX</code> sözlüğüne yeni eşleme ekleyin.<br>• <b>TC boş görünüyor:</b> normaldir, öğrenci işaretlememiştir.<br>• <b>Telefon cevaba karışıyor:</b> sistem 10/11 haneli telefonu otomatik ayırır.</p></div></div>
        </div>""",unsafe_allow_html=True)

elif not anahtar_input:
    st.warning("⚠️ Lütfen sol panelden cevap anahtarını girin.")
elif not yuklenen:
    st.warning("⚠️ Lütfen sol panelden bir TXT dosyası yükleyin.")
else:
    try:
        raw=yuklenen.read(); metin=None
        for enc in ("utf-8","cp1254","latin-1"):
            try: metin=raw.decode(enc); break
            except UnicodeDecodeError: continue
        if metin is None:
            st.error("Dosya okunamadı."); st.stop()
    except Exception as e:
        st.error(f"Dosya okunamadı: {e}"); st.stop()

    with st.spinner("İşleniyor…"):
        sonuclar,hatalar=isle(metin,anahtar_input)
    if not sonuclar:
        st.error("Hiç öğrenci verisi işlenemedi. Dosya formatını kontrol edin."); st.stop()

    toplamlar=[r["toplam"] for r in sonuclar]; soru_sayisi=len(anahtar_input)
    ort=np.mean(toplamlar); std=np.std(toplamlar)
    gecme=sum(1 for t in toplamlar if t>=50); en_yuksek=max(toplamlar)

    st.markdown(f"""
    <div class="metric-grid">
        <div class="mc purple"><div class="mc-label">Öğrenci</div><div class="mc-val">{len(sonuclar)}</div><div class="mc-sub">{soru_sayisi} soru · {round(100/soru_sayisi,2)} pt/soru</div></div>
        <div class="mc green"><div class="mc-label">Sınıf Ortalaması</div><div class="mc-val">{ort:.1f}</div><div class="mc-sub">100 üzerinden</div></div>
        <div class="mc sky"><div class="mc-label">Standart Sapma</div><div class="mc-val">{std:.1f}</div><div class="mc-sub">Dağılım genişliği</div></div>
        <div class="mc amber"><div class="mc-label">Geçen (≥50)</div><div class="mc-val">{gecme}</div><div class="mc-sub">%{round(gecme/len(sonuclar)*100,1)} geçme oranı</div></div>
        <div class="mc rose"><div class="mc-label">En Yüksek</div><div class="mc-val">{en_yuksek:.1f}</div><div class="mc-sub">En düşük: {min(toplamlar):.1f}</div></div>
    </div>""",unsafe_allow_html=True)

    if hatalar:
        with st.expander(f"⚠️ {len(hatalar)} satırda sorun tespit edildi",expanded=False):
            for h in hatalar: st.markdown(f'<div class="warn-box">{h}</div>',unsafe_allow_html=True)

    st.markdown('<div class="sec-head">📈 Puan Dağılımı</div>',unsafe_allow_html=True)
    araliklar=[("0–24",0,24,"#f87171"),("25–49",25,49,"#fb923c"),("50–64",50,64,"#fbbf24"),
               ("65–79",65,79,"#34d399"),("80–89",80,89,"#22d3ee"),("90–100",90,100,"#7c6af7")]
    cnts=[sum(1 for t in toplamlar if lo<=t<=hi) for _,lo,hi,_ in araliklar]
    max_cnt=max(cnts) if cnts else 1
    dh='<div class="dist-bar-wrap">'
    for (lbl,lo,hi,color),cnt in zip(araliklar,cnts):
        pb=round(cnt/max_cnt*100) if max_cnt else 0
        dh+=f'<div class="dist-row"><div class="dist-lbl">{lbl}</div><div class="dist-bar-bg"><div class="dist-bar-fill" style="width:{pb}%;background:{color};"><span>{"%"+str(round(cnt/len(toplamlar)*100,1)) if cnt else ""}</span></div></div><div class="dist-cnt">{cnt}</div></div>'
    dh+='</div>'; st.markdown(dh,unsafe_allow_html=True)

    st.markdown('<div class="sec-head">🔍 Soru Analizi</div>',unsafe_allow_html=True)
    analiz=soru_analizi_hesapla(sonuclar,anahtar_input)
    st.markdown("""<div style="display:flex;gap:1.2rem;margin-bottom:0.7rem;font-size:0.75rem;color:#6b6880;">
        <span><span style="color:#34d399;font-weight:700;">■</span> Kolay ≥70%</span>
        <span><span style="color:#fbbf24;font-weight:700;">■</span> Orta 40-69%</span>
        <span><span style="color:#f87171;font-weight:700;">■</span> Zor &lt;40%</span></div>""",unsafe_allow_html=True)
    ch='<div class="qa-grid">'
    for a in analiz:
        cls="easy" if a["pct"]>=70 else ("mid" if a["pct"]>=40 else "hard")
        ch+=f'<div class="qa-card {cls}"><div class="q-no">S{a["soru"]}</div><div class="q-pct">%{a["pct"]:.0f}</div><div class="q-ans">{a["anahtar"]}</div></div>'
    ch+='</div>'; st.markdown(ch,unsafe_allow_html=True)

    with st.expander("📋 Soru Analizi Detayları",expanded=False):
        dfq=pd.DataFrame([{"Soru":f"S{a['soru']}","Anahtar":a["anahtar"],"Doğru":a["dogru"],
            "Yanlış":a["yanlis"],"Boş":a["bos"],"Doğru %":a["pct"],
            "Zorluk":"🟢 Kolay" if a["pct"]>=70 else ("🟡 Orta" if a["pct"]>=40 else "🔴 Zor")} for a in analiz])
        st.dataframe(dfq,use_container_width=True,hide_index=True,
                     column_config={"Doğru %":st.column_config.ProgressColumn("Doğru %",min_value=0,max_value=100,format="%.1f%%")})

    st.markdown('<div class="sec-head">📋 Önizleme Tablosu</div>',unsafe_allow_html=True)
    sirali_ui=sorted(sonuclar,key=lambda r:tr_sort_key(r["ad_soyad"]))
    df_oniz=pd.DataFrame([{"#":i,"Ad Soyad":r["ad_soyad"],"TC":r["tc"] if r["tc"] else "—",
        "Cep":r["cep"] if r["cep"] else "—","Toplam (100)":r["toplam"],
        "Durum":"✓ Geçti" if r["toplam"]>=50 else "✗ Kaldı"} for i,r in enumerate(sirali_ui,1)])
    st.dataframe(df_oniz,use_container_width=True,hide_index=True,column_config={
        "#":st.column_config.NumberColumn("#",width="small"),
        "Ad Soyad":st.column_config.TextColumn("Ad Soyad",width="large"),
        "TC":st.column_config.TextColumn("TC No",width="medium"),
        "Cep":st.column_config.TextColumn("Cep Tel",width="medium"),
        "Toplam (100)":st.column_config.ProgressColumn("Toplam Puan",min_value=0,max_value=100,format="%.2f"),
        "Durum":st.column_config.TextColumn("Durum",width="small")})

    st.markdown("---")
    with st.spinner("Excel ve PDF hazırlanıyor…"):
        excel_buf=excel_olustur(sonuclar,anahtar_input)
        ders_bilgisi={"ders_adi":st.session_state.get("ders_adi",""),"ders_kodu":st.session_state.get("ders_kodu",""),
                      "bolum":st.session_state.get("bolum",""),"donem":st.session_state.get("donem",""),
                      "akademik_yil":st.session_state.get("akademik_yil","")}
        pdf_buf=pdf_olustur(sonuclar,anahtar_input,dosya_adi=yuklenen.name.rsplit(".",1)[0],
                            ders_bilgisi=ders_bilgisi,hoca_adi=st.session_state.get("hoca_adi",""))
    cepli=sum(1 for r in sonuclar if r.get("cep")); tcli=sum(1 for r in sonuclar if r.get("tc"))
    st.markdown(f'<div class="ok-box">✅ Hazır · 6 sheet · {len(sonuclar)} öğrenci · {gecme} geçti · '
                f'{tcli} TC · {cepli} cep tespit edildi</div>',unsafe_allow_html=True)
    c1,c2,_=st.columns([2,2,3])
    with c1:
        st.download_button("⬇️  Excel İndir  (.xlsx)",data=excel_buf,
            file_name=yuklenen.name.rsplit(".",1)[0]+"_sonuclar.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
    with c2:
        st.download_button("⬇️  PDF İndir  (.pdf)",data=pdf_buf,
            file_name=yuklenen.name.rsplit(".",1)[0]+"_rapor.pdf",mime="application/pdf",use_container_width=True)

st.markdown("""
<div class="footer-bar">
    Optik Notlandırma Sistemi · TC + Telefon Formatı<br>
    <b>Öğr. Gör. Hatice Tekiş</b> tarafından geliştirilmiştir
</div>""",unsafe_allow_html=True)
